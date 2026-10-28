import fitz
import pytest
from docx import Document

from backend.app.services.parser import parse_document

RU_TEXT = "Корова отказывается от корма уже третий день, температура тела повышена."


def test_parse_txt(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text(RU_TEXT, encoding="utf-8")

    result = parse_document(str(path))

    assert result["text"] == RU_TEXT
    assert result["file_type"] == "txt"
    assert result["language"] == "ru"


def test_parse_docx_skips_empty_paragraphs(tmp_path):
    path = tmp_path / "note.docx"
    doc = Document()
    doc.add_paragraph("Первый абзац.")
    doc.add_paragraph("   ")
    doc.add_paragraph("Второй абзац.")
    doc.save(path)

    result = parse_document(str(path))

    assert result["text"] == "Первый абзац.\nВторой абзац."
    assert result["file_type"] == "docx"


def test_parse_pdf_builds_page_map(tmp_path):
    path = tmp_path / "note.pdf"
    pdf = fitz.open()
    for text in ("Page one text", "Page two text"):
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()

    result = parse_document(str(path))

    assert "Page one text" in result["text"]
    assert "Page two text" in result["text"]
    assert [page for _, _, page in result["page_map"]] == [1, 2]


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")

    with pytest.raises(ValueError):
        parse_document(str(path))
